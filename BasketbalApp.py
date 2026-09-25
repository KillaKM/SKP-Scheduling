import copy
from datetime import datetime
import io
import random
import re
import openpyxl
from openpyxl.cell.cell import MergedCell
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="SKP Schedule & Task Automation", layout="wide"
)

# --- WACHTWOORDBEVEILIGING ---
def check_password():
    def password_entered():
        # Pas hier eventueel je eigen wachtwoord aan
        if st.session_state["password"] == "Tantalus2027!":
            st.session_state["password_correct"] = True
            del st.session_state["password"]
        else:
            st.session_state["password_correct"] = False

    if "password_correct" not in st.session_state:
        st.title("🔒 Inloggen vereist")
        st.info("Voer het wachtwoord in om toegang te krijgen tot het SKP Taaksysteem.")
        st.text_input("Wachtwoord:", type="password", on_change=password_entered, key="password")
        return False
    elif not st.session_state["password_correct"]:
        st.title("🔒 Inloggen vereist")
        st.text_input("Wachtwoord:", type="password", on_change=password_entered, key="password")
        st.error("Onjuist wachtwoord. Probeer het opnieuw.")
        return False
    return True

if not check_password():
    st.stop()
# ------------------------------

st.title("🏀 SKP Taakindeling & Scheidsrechters Systeem")

uploaded_file = st.sidebar.file_uploader(
    "Upload je Excel-bestand ('SKP Schedule 26_27 (BOARD).xlsx')", type=["xlsx"]
)

LOCK_COLS = ["Lock Ref 1", "Lock Ref 2", "Lock Scorer", "Lock Timer", "Lock 24s"]
TASK_COLS = ["Referee 1", "Referee 2", "Scorer", "Timer", "24 sec operator"]
LOCK_MAP = dict(zip(TASK_COLS, LOCK_COLS))


def ensure_lock_columns(df):
    df_res = df.copy()
    for col in LOCK_COLS:
        if col not in df_res.columns:
            df_res[col] = False
        else:
            df_res[col] = df_res[col].fillna(False).astype(bool)
    return df_res


def make_arrow_compatible(df):
    df_clean = df.copy()
    for col in df_clean.columns:
        if col in LOCK_COLS:
            df_clean[col] = df_clean[col].fillna(False).astype(bool)
        elif df_clean[col].dtype == object or str(col) in [
            "Referee 1", "Referee 2", "Scorer", "Timer", 
            "24 sec operator", "Home Team", "Away Team", "Division"
        ]:
            df_clean[col] = df_clean[col].fillna("").astype(str)
            df_clean[col] = df_clean[col].replace(
                {"nan": "", "None": "", "NoneType": ""}
            )
    return df_clean


def find_sheet(sheets_dict, candidates):
    for cand in candidates:
        for name in sheets_dict.keys():
            if cand.lower() == str(name).strip().lower():
                return name
    for cand in candidates:
        for name in sheets_dict.keys():
            if cand.lower() in str(name).strip().lower():
                return name
    return None


def find_col(df, candidates, fallback_index=None):
    for col in df.columns:
        c_clean = str(col).strip().lower().replace("_", " ").replace("-", " ")
        for cand in candidates:
            if cand.lower() == c_clean:
                return col
    for col in df.columns:
        c_clean = str(col).strip().lower()
        for cand in candidates:
            if cand.lower() in c_clean:
                return col
    if fallback_index is not None and len(df.columns) > fallback_index:
        return df.columns[fallback_index]
    return None


def clean_team_code(team_name):
    t = str(team_name).lower().replace("tantalus", "").replace("-", " ").strip()
    return "".join(t.split())


def build_division_map(divisions_df):
    div_map = {}
    if divisions_df is None or divisions_df.empty:
        return div_map
    df = divisions_df.copy()
    df.columns = [str(c).strip() for c in df.columns]

    team_col = find_col(df, ["team", "tantalus team", "teams", "teamnaam"], fallback_index=0)
    div_col = find_col(df, ["division", "divisie", "klasse", "poule"], fallback_index=1 if len(df.columns) > 1 else 0)

    for _, row in df.dropna(subset=[team_col]).iterrows():
        t_val = str(row[team_col]).strip()
        d_val = str(row[div_col]).strip()
        num = 5
        for d in ["1", "2", "3", "4", "5"]:
            if d in d_val.lower():
                num = int(d)
                break
        div_map[clean_team_code(t_val)] = num
    return div_map


def determine_division_for_team(team_str, div_map):
    if not team_str or str(team_str).strip() in ["", "nan", "None"]:
        return 5
    code = clean_team_code(team_str)
    if code in div_map:
        return div_map[code]
    for k, v in div_map.items():
        if k and (k in code or code in k):
            return v
    m = re.search(r'(?:mse|vse|xse|u\d+)[\s\-]*([1-5])', str(team_str).lower())
    if m:
        return int(m.group(1))
    return 5


def standardize_players_df(df):
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]

    col_last = find_col(df, ["last name", "lastname", "achternaam", "last"], fallback_index=1 if len(df.columns) > 1 else None)
    col_first = find_col(df, ["first name", "firstname", "voornaam", "first"], fallback_index=0 if len(df.columns) > 0 else None)
    col_dip = find_col(df, ["diploma", "licentie", "certificaat"], fallback_index=2 if len(df.columns) > 2 else None)
    col_comm = find_col(df, ["committee", "commissie"])
    col_extra_field = find_col(df, ["extra", "opmerking", "status", "notes"])
    col_season = find_col(df, ["full/ half season", "full/half season", "season", "seizoen", "half season", "half"])
    col_extra_pts = find_col(df, ["extra points", "extra punten", "commissie punten", "comm points"])

    renames = {}
    if col_first and col_first != "First name":
        renames[col_first] = "First name"
    if col_last and col_last != "Last name":
        renames[col_last] = "Last name"
    if col_dip and col_dip != "Diploma":
        renames[col_dip] = "Diploma"
    if col_comm and col_comm != "Committee":
        renames[col_comm] = "Committee"
    if col_extra_field and col_extra_field != "Extra":
        renames[col_extra_field] = "Extra"
    if col_season and col_season != "Full/ half season":
        renames[col_season] = "Full/ half season"
    if col_extra_pts and col_extra_pts != "Extra points":
        renames[col_extra_pts] = "Extra points"

    if renames:
        df = df.rename(columns=renames)

    if "First name" not in df.columns and len(df.columns) > 0:
        df["First name"] = df.iloc[:, 0]
    if "Last name" not in df.columns and len(df.columns) > 1:
        df["Last name"] = df.iloc[:, 1]
    if "Diploma" not in df.columns:
        df["Diploma"] = ""
    if "Committee" not in df.columns:
        df["Committee"] = ""
    if "Extra" not in df.columns:
        df["Extra"] = ""
    if "Full/ half season" not in df.columns:
        df["Full/ half season"] = "Full season"
    if "Extra points" not in df.columns:
        df["Extra points"] = 0.0
    else:
        df["Extra points"] = pd.to_numeric(df["Extra points"], errors="coerce").fillna(0.0)

    if "Total points" not in df.columns:
        df["Total points"] = df["Extra points"]
    else:
        df["Total points"] = pd.to_numeric(df["Total points"], errors="coerce").fillna(df["Extra points"])

    if "Referee" not in df.columns:
        df["Referee"] = 0
    if "Table duty" not in df.columns:
        df["Table duty"] = 0

    return df


def standardize_skp_df(df, div_map=None):
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]

    col_home = find_col(df, ["home team", "thuis team", "home", "thuis"])
    col_away = find_col(df, ["away team", "uit team", "away", "uit"])
    col_date = find_col(df, ["date", "datum"])
    col_time = find_col(df, ["time", "tijd"])
    col_div = find_col(df, ["division", "divisie", "poule", "klasse", "league", "div"])

    renames = {}
    if col_home and col_home != "Home Team":
        renames[col_home] = "Home Team"
    if col_away and col_away != "Away Team":
        renames[col_away] = "Away Team"
    if col_date and col_date != "Date":
        renames[col_date] = "Date"
    if col_time and col_time != "Time":
        renames[col_time] = "Time"
    if col_div and col_div != "Division":
        renames[col_div] = "Division"

    if renames:
        df = df.rename(columns=renames)

    if div_map is not None:
        div_vals = []
        for _, r in df.iterrows():
            h_team = str(r.get("Home Team", ""))
            d_num = determine_division_for_team(h_team, div_map)
            div_vals.append(f"Division {d_num}")
        df["Division"] = div_vals
    elif "Division" not in df.columns:
        df["Division"] = "Division 5"

    df = ensure_lock_columns(df)
    return df


def normalize_time_str(val):
    if pd.isna(val) or val is None:
        return ""
    s = str(val).strip()
    if len(s) >= 5 and ":" in s:
        return s[:5]
    return s


def normalize_date_str(val):
    if pd.isna(val) or val is None:
        return ""
    s = str(val).strip()
    if " " in s:
        s = s.split(" ")[0]
    return s


def parse_date_obj(val):
    if pd.isna(val) or val is None:
        return None
    if isinstance(val, (datetime, pd.Timestamp)):
        return val.date()
    s = normalize_date_str(val)
    for fmt in ["%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%y"]:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    m = re.search(r'(\d{1,2})[-/](\d{1,2})[-/](\d{2,4})', s)
    if m:
        d, m_val, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        try:
            return datetime(y, m_val, d).date()
        except ValueError:
            pass
    return None


def is_player_eligible_for_date(season_type, date_val):
    if not season_type:
        return True
    st_clean = str(season_type).strip().lower()
    if "full" in st_clean:
        return True

    d_obj = parse_date_obj(date_val)
    if not d_obj:
        return True

    pivot_first_half = datetime(2027, 1, 31).date()
    pivot_second_half = datetime(2027, 2, 1).date()

    if "first" in st_clean or "1st" in st_clean:
        return d_obj <= pivot_first_half
    elif "second" in st_clean or "2nd" in st_clean:
        return d_obj >= pivot_second_half

    return True


def build_team_busy_slots(all_games_df):
    busy_slots = set()
    if all_games_df is not None and not all_games_df.empty:
        for _, row in all_games_df.iterrows():
            d_val = normalize_date_str(row.get("Date"))
            t_val = normalize_time_str(row.get("Time"))
            h_team = str(row.get("Home Team", "")).lower()
            a_team = str(row.get("Away Team", "")).lower()

            if "tantalus" in h_team:
                c_h = clean_team_code(h_team)
                if c_h:
                    busy_slots.add((c_h, d_val, t_val))
            if "tantalus" in a_team:
                c_a = clean_team_code(a_team)
                if c_a:
                    busy_slots.add((c_a, d_val, t_val))
    return busy_slots


def is_player_playing(player_team, date_val, time_val, curr_home, curr_away, busy_slots):
    if not player_team or str(player_team).strip() in ["", "nan", "None"]:
        return False
    p_code = clean_team_code(player_team)
    if not p_code:
        return False

    h_str = str(curr_home).lower()
    a_str = str(curr_away).lower()
    if "tantalus" in h_str and p_code in clean_team_code(h_str):
        return True
    if "tantalus" in a_str and p_code in clean_team_code(a_str):
        return True

    d_norm = normalize_date_str(date_val)
    t_norm = normalize_time_str(time_val)

    if (p_code, d_norm, t_norm) in busy_slots:
        return True
    return False


def calculate_comm_points(comm_str, comm_points_map):
    if not comm_str or pd.isna(comm_str) or str(comm_str).strip() in ["", "nan", "None"]:
        return 0.0
    comms = [c.strip() for c in str(comm_str).split(",") if c.strip()]
    return sum(comm_points_map.get(c, 0.0) for c in comms)


def update_player_stats(sheets_dict):
    p_k = find_sheet(sheets_dict, ["Players skp", "Players", "Spelers"])
    s_k = find_sheet(sheets_dict, ["SKP", "Rooster"])
    c_k = find_sheet(sheets_dict, ["Committees", "Commissies"])

    if not p_k or not s_k or p_k not in sheets_dict or s_k not in sheets_dict:
        return sheets_dict

    players_df = standardize_players_df(sheets_dict[p_k].copy())
    skp_df = sheets_dict[s_k]
    committees_df = sheets_dict.get(c_k, pd.DataFrame())

    comm_points_map = {}
    if not committees_df.empty:
        col_comm = committees_df.columns[0]
        col_pts = committees_df.columns[-1] if len(committees_df.columns) > 1 else None
        for _, row in committees_df.dropna(subset=[col_comm]).iterrows():
            c_name = str(row[col_comm]).strip()
            try:
                pts = float(row[col_pts])
            except:
                pts = 0.0
            comm_points_map[c_name] = pts

    ref_counts = {}
    table_counts = {}

    for _, row in skp_df.iterrows():
        for ref_col in ["Referee 1", "Referee 2"]:
            val = str(row.get(ref_col, "")).strip()
            if val and val not in ["nan", "None", "x", ""]:
                ref_counts[val] = ref_counts.get(val, 0) + 1

        for duty_col in ["Scorer", "Timer", "24 sec operator"]:
            val = str(row.get(duty_col, "")).strip()
            if val and val not in ["nan", "None", "x", ""]:
                table_counts[val] = table_counts.get(val, 0) + 1

    for idx, row in players_df.iterrows():
        f_name = str(row.get("First name", "")).strip()
        l_name = str(row.get("Last name", "")).strip()

        if not l_name or l_name in ["nan", "", "none", "None"]:
            continue

        full_name = f"{f_name} {l_name}".strip()
        r_count = ref_counts.get(full_name, 0)
        t_count = table_counts.get(full_name, 0)

        comm_name = str(row.get("Committee", "")).strip()
        if comm_name:
            extra_pts = calculate_comm_points(comm_name, comm_points_map)
        else:
            try:
                extra_pts = float(row.get("Extra points", 0.0))
            except:
                extra_pts = 0.0

        total_pts = extra_pts + (r_count * 2) + (t_count * 1)

        players_df.at[idx, "Extra points"] = extra_pts
        players_df.at[idx, "Referee"] = r_count
        players_df.at[idx, "Table duty"] = t_count
        players_df.at[idx, "Total points"] = total_pts

    sheets_dict[p_k] = players_df
    return sheets_dict


def get_team_player_groups(players_df):
    team_groups = {}
    current_team = "Overig / Geen Team"
    for idx, row in players_df.iterrows():
        f_val = str(row.get("First name", "")).strip()
        l_val = str(row.get("Last name", "")).strip()
        if l_val in ["nan", "", "none", "None"]:
            if f_val and f_val not in ["nan", "", "none", "None"]:
                current_team = f_val
                if current_team not in team_groups:
                    team_groups[current_team] = []
        else:
            if current_team not in team_groups:
                team_groups[current_team] = []
            team_groups[current_team].append(idx)
    return team_groups


def fill_vacated_tasks(skp_df, removed_player_name, valid_pool, busy_slots, tantalus_div_map):
    task_cols = ["Referee 1", "Referee 2", "Scorer", "Timer", "24 sec operator"]
    replaced_count = 0

    for s_idx in skp_df.index:
        m_date = str(skp_df.at[s_idx, "Date"])
        m_time = skp_df.at[s_idx, "Time"]
        d_norm = normalize_date_str(m_date)
        t_norm = normalize_time_str(m_time)
        h_team = str(skp_df.at[s_idx, "Home Team"])
        a_team = str(skp_df.at[s_idx, "Away Team"])
        div_num = determine_division_for_team(h_team, tantalus_div_map)

        for t_col in task_cols:
            if str(skp_df.at[s_idx, t_col]).strip() == removed_player_name:
                skp_df.at[s_idx, t_col] = ""
                is_ref = "Referee" in t_col
                candidates = []

                for cand_name, cand in valid_pool.items():
                    if cand_name == removed_player_name:
                        continue
                    if is_ref and cand["Diploma_Rank"] == 0:
                        continue
                    if not is_ref and (cand["Diploma_Rank"] > 0 or cand.get("is_assistant_coach", False)):
                        continue

                    if is_ref and cand["Diploma"] == "BS3" and div_num != 2:
                        continue

                    if is_ref and "aurelie" in cand_name.lower():
                        if not ("mse" in h_team.lower() or "mse" in a_team.lower()) or div_num != 2:
                            continue

                    if not is_player_eligible_for_date(cand["Season_Type"], m_date):
                        continue

                    if is_player_playing(cand["Team"], m_date, m_time, h_team, a_team, busy_slots):
                        continue

                    already_busy = False
                    for _, check_row in skp_df.iterrows():
                        if normalize_date_str(check_row.get("Date")) == d_norm and normalize_time_str(check_row.get("Time")) == t_norm:
                            if any(str(check_row.get(c, "")).strip() == cand_name for c in task_cols):
                                already_busy = True
                                break
                    if already_busy:
                        continue

                    priority = 1
                    if is_ref:
                        dip = cand["Diploma"]
                        if div_num in [4, 5] and dip in ["BS1", "BS2"]:
                            priority = 10
                        elif div_num == 3 and dip in ["BS2", "L3", "L4"]:
                            priority = 10
                        elif div_num <= 2 and dip in ["BS3", "L3", "L4"]:
                            priority = 10

                    asst_penalty = 1 if cand.get("is_assistant_coach", False) else 0
                    candidates.append((asst_penalty, cand["Extra points"], priority, cand_name))

                candidates.sort(key=lambda x: (x[0], x[1], -x[2]))
                if candidates:
                    replacement = candidates[0][3]
                    skp_df.at[s_idx, t_col] = replacement
                    replaced_count += 1

    return skp_df, replaced_count


if uploaded_file is not None:
    if "original_sheets" not in st.session_state or st.sidebar.button("🔄 Bestand opnieuw inlezen"):
        file_bytes = uploaded_file.getvalue()
        st.session_state["file_bytes"] = file_bytes

        xls = pd.ExcelFile(io.BytesIO(file_bytes))
        raw_sheets = {}
        for sheet in xls.sheet_names:
            df = xls.parse(sheet)
            raw_sheets[sheet] = df.loc[:, ~df.columns.astype(str).str.contains("^Unnamed")]

        div_sheet_name = find_sheet(raw_sheets, ["Divisions", "Divisies"])
        div_map = build_division_map(raw_sheets[div_sheet_name]) if div_sheet_name else {}

        orig_sheets = {}
        for sheet, df in raw_sheets.items():
            if "player" in sheet.lower():
                df = standardize_players_df(df)
            elif "skp" in sheet.lower() and "player" not in sheet.lower():
                df = standardize_skp_df(df, div_map=div_map)
                task_cols = ["Referee 1", "Referee 2", "Scorer", "Timer", "24 sec operator"]
                for col in task_cols:
                    if col in df.columns:
                        df[col] = df[col].fillna("").astype(str).replace({"nan": "", "None": ""})
            orig_sheets[sheet] = df

        st.session_state["original_sheets"] = orig_sheets
        st.session_state["sheets"] = copy.deepcopy(orig_sheets)
        st.session_state["changelog"] = []
        st.session_state["indeling_gedaan"] = False
        st.session_state["wb_original"] = openpyxl.load_workbook(io.BytesIO(file_bytes))

if "sheets" in st.session_state:
    sheets = st.session_state["sheets"]

    players_key = find_sheet(sheets, ["Players skp", "Players", "Spelers"]) or "Players"
    skp_key = find_sheet(sheets, ["SKP", "Rooster"]) or "SKP"
    comm_key = find_sheet(sheets, ["Committees", "Commissies"]) or "Committees"
    all_games_key = find_sheet(sheets, ["All games", "all games", "ALL GAMES"]) or "all games"
    div_key = find_sheet(sheets, ["Divisions", "Divisies"]) or "Divisions"

    divisions_df = sheets.get(div_key, pd.DataFrame())
    tantalus_div_map = build_division_map(divisions_df)

    if skp_key in sheets:
        sheets[skp_key] = standardize_skp_df(sheets[skp_key], div_map=tantalus_div_map)

    if "action_feedback" in st.session_state and st.session_state["action_feedback"]:
        msg_type, msg_text = st.session_state["action_feedback"]
        if msg_type == "success":
            st.success(msg_text)
        elif msg_type == "warning":
            st.warning(msg_text)
        elif msg_type == "info":
            st.info(msg_text)
        st.session_state["action_feedback"] = None

    # --- SELECTIE SHEET BOVEN HET HOOFDVENSTER ---
    tab_names = list(sheets.keys())
    col_sel_sheet, _ = st.columns([1, 2])
    with col_sel_sheet:
        selected_tab = st.selectbox(
            "📋 Kies een Sheet om te bekijken/bewerken:",
            tab_names,
            key="sheet_selector_main"
        )

    st.subheader(f"Sheet: {selected_tab}")
    display_df = make_arrow_compatible(sheets[selected_tab])

    column_config = {}
    if selected_tab == skp_key:
        for l_col in LOCK_COLS:
            column_config[l_col] = st.column_config.CheckboxColumn(
                f"🔒 {l_col.replace('Lock ', '')}",
                help="Vink aan om deze taak vast te zetten (beschermd tegen overschrijven en wissen).",
                default=False,
            )

    edited_df = st.data_editor(
        display_df,
        num_rows="dynamic",
        width="stretch",
        key=f"editor_{selected_tab}",
        column_config=column_config
    )

    if not edited_df.equals(sheets[selected_tab]):
        sheets[selected_tab] = edited_df
        sheets = update_player_stats(sheets)
        st.rerun()

    # --- 1. MENU: LEDENBEHEER PER TEAM ---
    st.sidebar.divider()
    with st.sidebar.expander("👥 Ledenbeheer per Team", expanded=False):
        if players_key in sheets:
            players_manage_df = standardize_players_df(sheets[players_key])
            team_groups_manage = get_team_player_groups(players_manage_df)
            team_names_list = [t for t in team_groups_manage.keys() if t != "Overig / Geen Team"]
            if not team_names_list:
                team_names_list = list(team_groups_manage.keys())

            action_member = st.radio("Actie:", ["➕ Lid Toevoegen", "✏️ Lid Wijzigen", "🗑️ Lid Verwijderen"])

            if action_member == "➕ Lid Toevoegen":
                target_team = st.selectbox("Selecteer Team:", team_names_list, key="add_mem_team")
                new_first_name = st.text_input("Voornaam:", key="add_mem_fn")
                new_last_name = st.text_input("Achternaam:", key="add_mem_ln")
                new_dip = st.selectbox("Diploma:", ["", "BS1", "BS2", "BS3", "L3", "L4"], key="add_mem_dip")
                new_season = st.selectbox("Seizoen:", ["Full season", "First half", "Second half"], key="add_mem_season")

                if st.button("➕ Voeg Lid Toe Aan Team"):
                    if not new_first_name or not new_last_name:
                        st.error("Vul voornaam en achternaam in.")
                    else:
                        target_indices = team_groups_manage.get(target_team, [])
                        new_row_data = {
                            "First name": new_first_name.strip(),
                            "Last name": new_last_name.strip(),
                            "Diploma": new_dip,
                            "Full/ half season": new_season,
                            "Extra points": 0.0,
                            "Referee": 0,
                            "Table duty": 0,
                            "Total points": 0.0,
                            "Extra": "",
                            "Committee": ""
                        }

                        if target_indices:
                            insert_at = max(target_indices) + 1
                        else:
                            insert_at = len(players_manage_df)

                        df_top = players_manage_df.iloc[:insert_at]
                        df_bottom = players_manage_df.iloc[insert_at:]
                        new_member_df = pd.DataFrame([new_row_data])
                        updated_players_res = pd.concat([df_top, new_member_df, df_bottom], ignore_index=True)

                        sheets[players_key] = updated_players_res
                        sheets = update_player_stats(sheets)
                        st.session_state["changelog"].append({
                            "Actie": "Lid Toegevoegd",
                            "Details": f"{new_first_name} {new_last_name} toegevoegd aan {target_team}.",
                        })
                        st.session_state["action_feedback"] = (
                            "success",
                            f"Lid {new_first_name} {new_last_name} is succesvol toegevoegd aan {target_team}!"
                        )
                        st.rerun()

            elif action_member == "✏️ Lid Wijzigen":
                edit_team = st.selectbox("Selecteer Team:", team_names_list, key="edit_mem_team")
                edit_indices = team_groups_manage.get(edit_team, [])
                if edit_indices:
                    edit_player_dict = {}
                    for idx_m in edit_indices:
                        fn = str(players_manage_df.at[idx_m, "First name"]).strip()
                        ln = str(players_manage_df.at[idx_m, "Last name"]).strip()
                        edit_player_dict[f"{fn} {ln}"] = idx_m

                    selected_edit_name = st.selectbox("Selecteer lid om te bewerken:", list(edit_player_dict.keys()), key="edit_mem_name")
                    idx_to_edit = edit_player_dict[selected_edit_name]

                    curr_fn = str(players_manage_df.at[idx_to_edit, "First name"])
                    curr_ln = str(players_manage_df.at[idx_to_edit, "Last name"])
                    curr_dip = str(players_manage_df.at[idx_to_edit, "Diploma"]).strip().upper()
                    curr_season = str(players_manage_df.at[idx_to_edit, "Full/ half season"]).strip()

                    dip_options = ["", "BS1", "BS2", "BS3", "L3", "L4"]
                    dip_idx = dip_options.index(curr_dip) if curr_dip in dip_options else 0

                    season_options = ["Full season", "First half", "Second half"]
                    season_idx = 0
                    for s_i, s_opt in enumerate(season_options):
                        if s_opt.lower() in curr_season.lower():
                            season_idx = s_i
                            break

                    up_fn = st.text_input("Voornaam:", value=curr_fn, key="up_mem_fn")
                    up_ln = st.text_input("Achternaam:", value=curr_ln, key="up_mem_ln")
                    up_dip = st.selectbox("Diploma:", dip_options, index=dip_idx, key="up_mem_dip")
                    up_season = st.selectbox("Seizoen:", season_options, index=season_idx, key="up_mem_season")

                    if st.button("💾 Sla Wijzigingen Lid Op"):
                        players_manage_df.at[idx_to_edit, "First name"] = up_fn.strip()
                        players_manage_df.at[idx_to_edit, "Last name"] = up_ln.strip()
                        players_manage_df.at[idx_to_edit, "Diploma"] = up_dip
                        players_manage_df.at[idx_to_edit, "Full/ half season"] = up_season

                        sheets[players_key] = players_manage_df
                        sheets = update_player_stats(sheets)
                        st.session_state["changelog"].append({
                            "Actie": "Lid Gewijzigd",
                            "Details": f"{up_fn} {up_ln} ({edit_team}): diploma={up_dip}, seizoen={up_season}.",
                        })
                        st.session_state["action_feedback"] = (
                            "success",
                            f"Gegevens voor {up_fn} {up_ln} zijn succesvol bijgewerkt!"
                        )
                        st.rerun()
                else:
                    st.info("Geen leden gevonden in dit team.")

            else:
                del_team = st.selectbox("Selecteer Team:", team_names_list, key="del_mem_team")
                del_indices = team_groups_manage.get(del_team, [])
                if del_indices:
                    del_player_dict = {}
                    for idx_m in del_indices:
                        fn = str(players_manage_df.at[idx_m, "First name"]).strip()
                        ln = str(players_manage_df.at[idx_m, "Last name"]).strip()
                        del_player_dict[f"{fn} {ln}"] = idx_m

                    selected_del_name = st.selectbox("Selecteer lid om te verwijderen:", list(del_player_dict.keys()), key="del_mem_name")
                    if st.button("🗑️ Verwijder Dit Lid Definitief"):
                        idx_to_del = del_player_dict[selected_del_name]
                        updated_players_res = players_manage_df.drop(index=idx_to_del).reset_index(drop=True)
                        sheets[players_key] = updated_players_res

                        skp_df_del = sheets.get(skp_key, pd.DataFrame())
                        all_games_df = sheets.get(all_games_key, pd.DataFrame())
                        if not all_games_df.empty:
                            all_games_df = standardize_skp_df(all_games_df)
                        busy_slots = build_team_busy_slots(all_games_df)

                        player_team_map = {}
                        curr_t = ""
                        for _, p_row in updated_players_res.iterrows():
                            f_val = str(p_row.get("First name", "")).strip()
                            l_val = str(p_row.get("Last name", "")).strip()
                            if l_val in ["nan", "", "none", "None"]:
                                if f_val and f_val not in ["nan", "", "none", "None"]:
                                    curr_t = f_val
                            else:
                                player_team_map[f"{f_val} {l_val}".strip()] = curr_t

                        diploma_ranking = {
                            "L4": 5, "L3": 4, "BS3": 3, "BS2": 2, "BS1": 1, "nan": 0, "": 0
                        }

                        valid_pool = {}
                        for idx_p, p_row in updated_players_res.iterrows():
                            l_val = str(p_row.get("Last name", "")).strip()
                            f_val = str(p_row.get("First name", "")).strip()
                            extra_field_val = str(p_row.get("Extra", "")).strip().lower()
                            comm_field_val = str(p_row.get("Committee", "")).strip().lower()
                            combined_roles = f"{extra_field_val} {comm_field_val}"

                            if l_val in ["nan", "", "none", "None"] or "recreational for now" in combined_roles:
                                continue

                            is_assistant_coach = "assistant coach" in combined_roles
                            is_board = "board" in combined_roles or "bestuur" in combined_roles
                            is_head_coach = ("coach" in combined_roles) and not is_assistant_coach

                            if is_board or is_head_coach:
                                continue

                            full_name = f"{f_val} {l_val}".strip()
                            d_val = str(p_row.get("Diploma", "")).strip().upper().replace(" ", "").replace("-", "")
                            rank = diploma_ranking.get(d_val, 0)
                            team = player_team_map.get(full_name, "")
                            season_type = str(p_row.get("Full/ half season", "Full season")).strip()
                            try:
                                extra_p = float(p_row.get("Extra points", 0.0))
                            except:
                                extra_p = 0.0

                            valid_pool[full_name] = {
                                "Full Name": full_name,
                                "Team": team,
                                "Diploma": d_val,
                                "Diploma_Rank": rank,
                                "Extra points": extra_p,
                                "Season_Type": season_type,
                                "is_assistant_coach": is_assistant_coach,
                            }

                        skp_df_del, num_replaced = fill_vacated_tasks(
                            skp_df_del, selected_del_name, valid_pool, busy_slots, tantalus_div_map
                        )
                        sheets[skp_key] = make_arrow_compatible(skp_df_del)
                        sheets = update_player_stats(sheets)

                        st.session_state["changelog"].append({
                            "Actie": "Lid Verwijderd & Rooster Hersteld",
                            "Details": f"{selected_del_name} verwijderd uit {del_team}. {num_replaced} taken direct herverdeeld naar beschikbare leden.",
                        })
                        st.session_state["action_feedback"] = (
                            "success",
                            f"{selected_del_name} is definitief verwijderd. {num_replaced} openstaande taken zijn direct herverdeeld!"
                        )
                        st.rerun()
                else:
                    st.info("Geen leden gevonden in dit team.")

    # --- 2. SAMENGEVOEGD MENU: ROOSTER INDELEN & BEHEER (PER DAG & PER REGEL) ---
    st.sidebar.divider()
    target_match_indices = []

    with st.sidebar.expander("📅 Rooster Indelen & Beheer", expanded=True):
        scope_mode = st.radio("Toepassen op:", ["Per dag(en)", "Per regel (specifieke wedstrijd)"], key="scope_mode")

        st.markdown("#### 1. Indeling Toepassen")
        if skp_key in sheets:
            skp_df_ctrl = sheets[skp_key]
            
            if scope_mode == "Per dag(en)":
                if "Date" in skp_df_ctrl.columns:
                    unique_dates = list(skp_df_ctrl["Date"].dropna().unique())

                    def toggle_all_days():
                        select_state = st.session_state.get("select_all_days", False)
                        for date_val in unique_dates:
                            st.session_state[f"date_chk_{date_val}"] = select_state

                    st.checkbox("✅ Selecteer Alle Dagen", key="select_all_days", on_change=toggle_all_days)

                    chosen_dates = []
                    for date_val in unique_dates:
                        chk_key = f"date_chk_{date_val}"
                        if chk_key not in st.session_state:
                            st.session_state[chk_key] = False
                        if st.checkbox(str(date_val), key=chk_key):
                            chosen_dates.append(date_val)

                    chosen_dates_norm = {normalize_date_str(d) for d in chosen_dates}
                    for idx_r in skp_df_ctrl.index:
                        if normalize_date_str(skp_df_ctrl.at[idx_r, "Date"]) in chosen_dates_norm:
                            target_match_indices.append(idx_r)
            else:
                row_choices = []
                for idx_r, r_val in skp_df_ctrl.iterrows():
                    d = str(r_val.get("Date", ""))
                    t = str(r_val.get("Time", ""))
                    h = str(r_val.get("Home Team", ""))
                    a = str(r_val.get("Away Team", ""))
                    row_choices.append((idx_r, f"Rij {idx_r + 1}: {d} ({t}) - {h} vs {a}"))
                sel_row_idx = st.selectbox("Kies wedstrijd:", [r[0] for r in row_choices], format_func=lambda x: dict(row_choices)[x], key="sel_row_assign")
                target_match_indices = [sel_row_idx]

            max_tasks_per_day = st.slider("Max. taken per speler per dag", 1, 4, 2)
            start_auto_btn = st.button("🤖 Start indeling voor selectie")

        st.divider()
        st.markdown("#### 2. Indeling Wissen")
        if scope_mode == "Per dag(en)":
            reset_mode = st.radio("Kies wis-modus:", ["Hele rooster wissen", "Indeling per dag wissen"], key="reset_mode_radio")

            if reset_mode == "Hele rooster wissen":
                if st.button("🗑️ Wis het hele rooster (excl. 🔒)", key="btn_clear_all"):
                    if skp_key in sheets:
                        skp_df_w = sheets[skp_key]
                        task_cols_to_clear = ["Referee 1", "Referee 2", "Scorer", "Timer", "24 sec operator"]
                        for t_c in task_cols_to_clear:
                            l_col = LOCK_MAP[t_c]
                            if t_c in skp_df_w.columns:
                                for i in skp_df_w.index:
                                    is_locked = bool(skp_df_w.at[i, l_col]) if l_col in skp_df_w.columns else False
                                    if not is_locked:
                                        skp_df_w.at[i, t_c] = ""
                        sheets[skp_key] = make_arrow_compatible(skp_df_w)
                        sheets = update_player_stats(sheets)
                        st.session_state["changelog"].append({"Actie": "Rooster Wissen", "Details": "Hele rooster gewist (met behoud van vastgezette taken)."})
                        st.session_state["action_feedback"] = ("info", "Het volledige rooster is succesvol gewist (vastgezette taken zijn behouden).")
                        st.rerun()
            else:
                if skp_key in sheets:
                    skp_df_temp = sheets[skp_key]
                    if "Date" in skp_df_temp.columns:
                        reset_dates_list = list(skp_df_temp["Date"].dropna().unique())
                        selected_reset_date = st.selectbox("Selecteer te wissen dag:", reset_dates_list, key="sel_reset_date")
                        if st.button(f"🗑️ Wis dag {selected_reset_date}", key="btn_clear_day"):
                            d_reset_norm = normalize_date_str(selected_reset_date)
                            mask_reset = skp_df_temp["Date"].apply(normalize_date_str) == d_reset_norm
                            task_cols_to_clear = ["Referee 1", "Referee 2", "Scorer", "Timer", "24 sec operator"]
                            for idx_r in skp_df_temp[mask_reset].index:
                                for t_c in task_cols_to_clear:
                                    l_col = LOCK_MAP[t_c]
                                    if t_c in skp_df_temp.columns:
                                        is_locked = bool(skp_df_temp.at[idx_r, l_col]) if l_col in skp_df_temp.columns else False
                                        if not is_locked:
                                            skp_df_temp.at[idx_r, t_c] = ""
                            sheets[skp_key] = make_arrow_compatible(skp_df_temp)
                            sheets = update_player_stats(sheets)
                            st.session_state["changelog"].append({"Actie": "Rooster Wissen", "Details": f"Indeling voor {selected_reset_date} gewist."})
                            st.session_state["action_feedback"] = ("info", f"De indeling voor {selected_reset_date} is succesvol gewist.")
                            st.rerun()
        else:
            if skp_key in sheets:
                sel_row_to_clear = st.selectbox("Kies wedstrijd om te wissen:", [r[0] for r in row_choices], format_func=lambda x: dict(row_choices)[x], key="sel_row_clear")
                if st.button(f"🗑️ Wis alleen Rij {sel_row_to_clear + 1}", key="btn_clear_row"):
                    skp_df_row = sheets[skp_key]
                    task_cols_to_clear = ["Referee 1", "Referee 2", "Scorer", "Timer", "24 sec operator"]
                    for t_c in task_cols_to_clear:
                        l_col = LOCK_MAP[t_c]
                        if t_c in skp_df_row.columns:
                            is_locked = bool(skp_df_row.at[sel_row_to_clear, l_col]) if l_col in skp_df_row.columns else False
                            if not is_locked:
                                skp_df_row.at[sel_row_to_clear, t_c] = ""
                    sheets[skp_key] = make_arrow_compatible(skp_df_row)
                    sheets = update_player_stats(sheets)
                    st.session_state["changelog"].append({"Actie": "Regel Gewist", "Details": f"Taken in rij {sel_row_to_clear + 1} gewist."})
                    st.session_state["action_feedback"] = ("info", f"De taken voor rij {sel_row_to_clear + 1} zijn gewist.")
                    st.rerun()

    # --- UITVOEREN VAN AUTOMATISCHE INDELING ---
    if "start_auto_btn" in locals() and start_auto_btn:
        if not target_match_indices:
            st.session_state["action_feedback"] = ("warning", "Geen wedstrijden geselecteerd om in te delen.")
            st.rerun()
        else:
            sheets = update_player_stats(sheets)
            skp_df = sheets[skp_key]
            all_games_df = sheets.get(all_games_key, pd.DataFrame())
            if not all_games_df.empty:
                all_games_df = standardize_skp_df(all_games_df)

            skp_df = standardize_skp_df(skp_df, div_map=tantalus_div_map)

            for col in ["Referee 1", "Referee 2", "Scorer", "Timer", "24 sec operator"]:
                if col in skp_df.columns:
                    skp_df[col] = skp_df[col].fillna("").astype(str).replace({"nan": "", "None": ""})

            players_df = standardize_players_df(sheets.get(players_key, pd.DataFrame()))

            player_team_map = {}
            current_team = ""
            for _, p_row in players_df.iterrows():
                f_val = str(p_row.get("First name", "")).strip()
                l_val = str(p_row.get("Last name", "")).strip()
                if l_val in ["nan", "", "none", "None"]:
                    if f_val and f_val not in ["nan", "", "none", "None"]:
                        current_team = f_val
                else:
                    full_p_name = f"{f_val} {l_val}".strip()
                    player_team_map[full_p_name] = current_team

            diploma_ranking = {
                "L4": 5, "L3": 4, "BS3": 3, "BS2": 2, "BS1": 1, "nan": 0, "": 0
            }

            valid_players_dict = {}
            for idx, p_row in players_df.iterrows():
                l_val = str(p_row.get("Last name", "")).strip()
                f_val = str(p_row.get("First name", "")).strip()
                if l_val not in ["nan", "", "none", "None"]:
                    extra_field_val = str(p_row.get("Extra", "")).strip().lower()
                    comm_field_val = str(p_row.get("Committee", "")).strip().lower()
                    combined_roles = f"{extra_field_val} {comm_field_val}"

                    if "recreational for now" in combined_roles:
                        continue

                    is_assistant_coach = "assistant coach" in combined_roles
                    is_board = "board" in combined_roles or "bestuur" in combined_roles
                    is_head_coach = ("coach" in combined_roles) and not is_assistant_coach

                    if is_board or is_head_coach:
                        continue

                    full_name = f"{f_val} {l_val}".strip()
                    d_val = str(p_row.get("Diploma", "")).strip().upper().replace(" ", "").replace("-", "")
                    rank = diploma_ranking.get(d_val, 0)
                    team = player_team_map.get(full_name, "")
                    season_type = str(p_row.get("Full/ half season", "Full season")).strip()

                    try:
                        extra = float(p_row.get("Extra points", 0.0))
                    except:
                        extra = 0.0

                    valid_players_dict[full_name] = {
                        "First name": f_val,
                        "Last name": l_val,
                        "Full Name": full_name,
                        "Team": team,
                        "Diploma": d_val,
                        "Diploma_Rank": rank,
                        "Extra points": extra,
                        "Season_Type": season_type,
                        "is_assistant_coach": is_assistant_coach,
                        "original_idx": idx
                    }

            # Maak niet-vergrendelde cellen in de doelrijen leeg
            for idx in target_match_indices:
                for col_c in ["Referee 1", "Referee 2", "Scorer", "Timer", "24 sec operator"]:
                    l_col = LOCK_MAP[col_c]
                    is_locked = bool(skp_df.at[idx, l_col]) if l_col in skp_df.columns else False
                    if col_c in skp_df.columns and not is_locked:
                        curr_v = str(skp_df.at[idx, col_c]).strip()
                        if curr_v.lower() != "x":
                            skp_df.at[idx, col_c] = ""

            busy_game_slots = build_team_busy_slots(all_games_df)

            ref_tasks_counter = {p: 0 for p in valid_players_dict}
            table_tasks_counter = {p: 0 for p in valid_players_dict}
            player_busy_times = {p: set() for p in valid_players_dict}
            player_day_task_counts = {p: {} for p in valid_players_dict}

            # Tel reeds bezette / vastgezette taken mee
            for _, row in skp_df.iterrows():
                d_val = normalize_date_str(row.get("Date"))
                t_val = normalize_time_str(row.get("Time"))

                for col in ["Referee 1", "Referee 2"]:
                    name = str(row.get(col, "")).strip()
                    if name in ref_tasks_counter:
                        ref_tasks_counter[name] += 1
                        player_busy_times[name].add((d_val, t_val))
                        player_day_task_counts[name][d_val] = player_day_task_counts[name].get(d_val, 0) + 1
                for col in ["Scorer", "Timer", "24 sec operator"]:
                    name = str(row.get(col, "")).strip()
                    if name in table_tasks_counter:
                        table_tasks_counter[name] += 1
                        player_busy_times[name].add((d_val, t_val))
                        player_day_task_counts[name][d_val] = player_day_task_counts[name].get(d_val, 0) + 1

            def get_row_div_val(r_idx):
                h_t = str(skp_df.at[r_idx, "Home Team"])
                return determine_division_for_team(h_t, tantalus_div_map)

            target_match_indices.sort(key=lambda idx_m: get_row_div_val(idx_m))

            for idx in target_match_indices:
                home_team = str(skp_df.at[idx, "Home Team"])
                away_team = str(skp_df.at[idx, "Away Team"])

                if "tantalus" not in home_team.lower():
                    continue

                m_date = str(skp_df.at[idx, "Date"])
                d_norm = normalize_date_str(m_date)
                m_time = skp_df.at[idx, "Time"]
                t_norm = normalize_time_str(m_time)
                div_num = determine_division_for_team(home_team, tantalus_div_map)
                skp_df.at[idx, "Division"] = f"Division {div_num}"

                assigned_in_match = {
                    str(skp_df.at[idx, c]).strip() for c in ["Referee 1", "Referee 2", "Scorer", "Timer", "24 sec operator"]
                    if str(skp_df.at[idx, c]).strip() not in ["", "nan", "None", "x"]
                }

                # --- SCHEIDSRECHTERS ---
                is_mse1 = "mse 1" in clean_team_code(home_team) or "mse1" in clean_team_code(home_team)
                if div_num != 1 and not is_mse1:
                    for ref_col in ["Referee 1", "Referee 2"]:
                        l_col = LOCK_MAP[ref_col]
                        is_locked = bool(skp_df.at[idx, l_col]) if l_col in skp_df.columns else False
                        if is_locked:
                            continue

                        curr_val = str(skp_df.at[idx, ref_col]).strip()
                        if curr_val.lower() == "x":
                            continue
                        if curr_val in ["", "None", "nan"]:
                            ref_candidates = []

                            for p_name, player in valid_players_dict.items():
                                rank = player["Diploma_Rank"]
                                if rank == 0:
                                    continue

                                diploma = player["Diploma"]

                                # Regel: BS3 mag alleen 2e divisie fluiten
                                if diploma == "BS3" and div_num != 2:
                                    continue

                                # Regel: Aurelie fluit alleen wedstrijden van mannen in 2e divisie
                                if "aurelie" in p_name.lower():
                                    if not ("mse" in home_team.lower() or "mse" in away_team.lower()) or div_num != 2:
                                        continue

                                is_asst = player.get("is_assistant_coach", False)
                                c_tasks = ref_tasks_counter[p_name]
                                if is_asst and c_tasks >= 1:
                                    continue

                                if not is_player_eligible_for_date(player["Season_Type"], m_date):
                                    continue

                                p_team = player["Team"]
                                total_pts = player["Extra points"] + (c_tasks * 2)

                                if (d_norm, t_norm) in player_busy_times[p_name]:
                                    continue
                                if is_player_playing(p_team, m_date, m_time, home_team, away_team, busy_game_slots):
                                    continue
                                if p_name in assigned_in_match:
                                    continue

                                day_count = player_day_task_counts[p_name].get(d_norm, 0)
                                over_day_penalty = 1 if day_count >= max_tasks_per_day else 0

                                eligible = False
                                priority = 0

                                if div_num == 5:
                                    if diploma in ["BS1", "BS2"]:
                                        eligible = True
                                        priority = 10 if diploma == "BS1" else 8
                                    elif rank >= 4:
                                        eligible = True
                                        priority = 2
                                elif div_num == 4:
                                    if diploma == "BS2":
                                        eligible = True
                                        priority = 10
                                    elif diploma in ["L3", "L4"]:
                                        eligible = True
                                        priority = 5
                                    elif diploma == "BS1":
                                        eligible = True
                                        priority = 2
                                elif div_num == 3:
                                    already_higher = any(valid_players_dict[a]["Diploma"] in ["L3", "L4"] for a in assigned_in_match if a in valid_players_dict)
                                    if not already_higher:
                                        if diploma in ["L3", "L4"]:
                                            eligible = True
                                            priority = 10
                                        elif diploma == "BS2":
                                            eligible = True
                                            priority = 6
                                        elif diploma == "BS1":
                                            eligible = True
                                            priority = 2
                                    else:
                                        if diploma == "BS2":
                                            eligible = True
                                            priority = 10
                                        elif diploma in ["L3", "L4"]:
                                            eligible = True
                                            priority = 6
                                        elif diploma == "BS1":
                                            eligible = True
                                            priority = 2
                                elif div_num == 2:
                                    already_l3 = any(valid_players_dict[a]["Diploma"] in ["L3", "L4"] for a in assigned_in_match if a in valid_players_dict)
                                    if not already_l3:
                                        if diploma in ["L3", "L4"]:
                                            eligible = True
                                            priority = 10
                                        elif diploma == "BS3":
                                            eligible = True
                                            priority = 8
                                        elif diploma == "BS2":
                                            eligible = True
                                            priority = 5
                                        elif diploma == "BS1":
                                            eligible = True
                                            priority = 1
                                    else:
                                        if diploma == "BS3":
                                            eligible = True
                                            priority = 10
                                        elif diploma in ["L3", "L4"]:
                                            eligible = True
                                            priority = 8
                                        elif diploma == "BS2":
                                            eligible = True
                                            priority = 5
                                        elif diploma == "BS1":
                                            eligible = True
                                            priority = 1

                                if eligible:
                                    over_cap_penalty = 1 if (total_pts >= 15.0 or player["Extra points"] >= 12.0) else 0
                                    asst_penalty = 1 if is_asst else 0

                                    ref_candidates.append({
                                        "name": p_name,
                                        "asst_coach": asst_penalty,
                                        "over_day": over_day_penalty,
                                        "over_cap": over_cap_penalty,
                                        "points": total_pts,
                                        "tasks": c_tasks,
                                        "priority": priority,
                                    })

                            ref_candidates.sort(key=lambda x: (
                                x["asst_coach"],
                                x["over_day"],
                                x["over_cap"],
                                x["points"],
                                x["tasks"],
                                -x["priority"],
                                random.random()
                            ))

                            if ref_candidates:
                                chosen = ref_candidates[0]["name"]
                                skp_df.at[idx, ref_col] = chosen
                                assigned_in_match.add(chosen)
                                player_busy_times[chosen].add((d_norm, t_norm))
                                player_day_task_counts[chosen][d_norm] = player_day_task_counts[chosen].get(d_norm, 0) + 1
                                ref_tasks_counter[chosen] += 1

                # --- TAFELTAKEN ---
                for col in ["Scorer", "Timer", "24 sec operator"]:
                    l_col = LOCK_MAP[col]
                    is_locked = bool(skp_df.at[idx, l_col]) if l_col in skp_df.columns else False
                    if is_locked:
                        continue

                    if col in skp_df.columns:
                        curr_val = str(skp_df.at[idx, col]).strip()
                        if curr_val in ["", "None", "nan"]:
                            duty_candidates = []

                            for p_name, player in valid_players_dict.items():
                                rank = player["Diploma_Rank"]
                                if rank > 0 or player.get("is_assistant_coach", False):
                                    continue

                                if not is_player_eligible_for_date(player["Season_Type"], m_date):
                                    continue

                                p_team = player["Team"]
                                c_tasks = table_tasks_counter[p_name]
                                total_pts = player["Extra points"] + c_tasks

                                if (d_norm, t_norm) in player_busy_times[p_name]:
                                    continue
                                if is_player_playing(p_team, m_date, m_time, home_team, away_team, busy_game_slots):
                                    continue
                                if p_name in assigned_in_match:
                                    continue

                                day_count = player_day_task_counts[p_name].get(d_norm, 0)
                                over_day_penalty = 1 if day_count >= max_tasks_per_day else 0
                                over_cap_penalty = 1 if (total_pts >= 15.0 or player["Extra points"] >= 12.0) else 0

                                duty_candidates.append({
                                    "name": p_name,
                                    "over_day": over_day_penalty,
                                    "over_cap": over_cap_penalty,
                                    "points": total_pts,
                                    "tasks": c_tasks,
                                })

                            duty_candidates.sort(key=lambda x: (
                                x["over_day"],
                                x["over_cap"],
                                x["points"],
                                x["tasks"],
                                random.random()
                            ))

                            if duty_candidates:
                                chosen = duty_candidates[0]["name"]
                                skp_df.at[idx, col] = chosen
                                assigned_in_match.add(chosen)
                                player_busy_times[chosen].add((d_norm, t_norm))
                                player_day_task_counts[chosen][d_norm] = player_day_task_counts[chosen].get(d_norm, 0) + 1
                                table_tasks_counter[chosen] += 1

            sheets[skp_key] = make_arrow_compatible(skp_df)
            sheets = update_player_stats(sheets)
            st.session_state["indeling_gedaan"] = True
            st.session_state["action_feedback"] = ("success", "Indeling succesvol uitgevoerd voor de geselecteerde wedstrijden!")
            st.rerun()

    # --- 3. MENU: COMMISSIES TOEWIJZEN ---
    st.sidebar.divider()
    if players_key in sheets and comm_key in sheets:
        comm_df = sheets[comm_key]
        col_comm = comm_df.columns[0]
        col_pts = comm_df.columns[-1] if len(comm_df.columns) > 1 else None

        comm_points_map_lookup = {}
        for _, c_row in comm_df.dropna(subset=[col_comm]).iterrows():
            try:
                pts = float(c_row[col_pts])
            except:
                pts = 0.0
            comm_points_map_lookup[str(c_row[col_comm]).strip()] = pts

        available_committees = [str(c).strip() for c in comm_df[col_comm].dropna().unique() if str(c).strip()]
        players_df = standardize_players_df(sheets[players_key])
        team_player_groups = get_team_player_groups(players_df)

        with st.sidebar.expander("🛠️ Open Commissie Menu (Per Team)", expanded=False):
            for team_name, player_indices in team_player_groups.items():
                if not player_indices:
                    continue

                with st.expander(f"🏀 {team_name} ({len(player_indices)} spelers)"):
                    for idx in player_indices:
                        f_name = str(players_df.at[idx, "First name"])
                        l_name = str(players_df.at[idx, "Last name"])
                        raw_comm = str(players_df.at[idx, "Committee"]).strip()

                        if raw_comm in ["nan", "None", "none", ""]:
                            current_comms_list = []
                        else:
                            current_comms_list = [c.strip() for c in raw_comm.split(",") if c.strip() in available_committees]

                        selected_comms = st.multiselect(
                            f"{f_name} {l_name}",
                            available_committees,
                            default=current_comms_list,
                            key=f"comm_multi_{idx}_{f_name}_{l_name}",
                        )
                        players_df.at[idx, "Committee"] = ", ".join(selected_comms)

            if st.button("💾 Sla Commissies op & Synchroniseer Rooster"):
                sheets[players_key] = players_df
                sheets = update_player_stats(sheets)

                skp_df_sync = sheets.get(skp_key, pd.DataFrame())
                updated_p_df = sheets[players_key]
                all_games_df = sheets.get(all_games_key, pd.DataFrame())
                if not all_games_df.empty:
                    all_games_df = standardize_skp_df(all_games_df)

                busy_slots = build_team_busy_slots(all_games_df)

                diploma_ranking = {
                    "L4": 5, "L3": 4, "BS3": 3, "BS2": 2, "BS1": 1, "nan": 0, "": 0
                }

                valid_pool = {}
                for idx, p_row in updated_p_df.iterrows():
                    l_val = str(p_row.get("Last name", "")).strip()
                    f_val = str(p_row.get("First name", "")).strip()
                    extra_field_val = str(p_row.get("Extra", "")).strip().lower()
                    comm_field_val = str(p_row.get("Committee", "")).strip().lower()
                    combined_roles = f"{extra_field_val} {comm_field_val}"

                    if l_val in ["nan", "", "none", "None"] or "recreational for now" in combined_roles:
                        continue

                    is_assistant_coach = "assistant coach" in combined_roles
                    is_board = "board" in combined_roles or "bestuur" in combined_roles
                    is_head_coach = ("coach" in combined_roles) and not is_assistant_coach

                    if is_board or is_head_coach:
                        continue

                    full_name = f"{f_val} {l_val}".strip()
                    d_val = str(p_row.get("Diploma", "")).strip().upper().replace(" ", "").replace("-", "")
                    rank = diploma_ranking.get(d_val, 0)
                    team = player_team_map.get(full_name, "")
                    season_type = str(p_row.get("Full/ half season", "Full season")).strip()
                    try:
                        extra_p = float(p_row.get("Extra points", 0.0))
                    except:
                        extra_p = 0.0

                    valid_pool[full_name] = {
                        "Full Name": full_name,
                        "Team": team,
                        "Diploma": d_val,
                        "Diploma_Rank": rank,
                        "Extra points": extra_p,
                        "Season_Type": season_type,
                        "is_assistant_coach": is_assistant_coach,
                    }

                relieved_actions = []
                task_cols = ["Referee 1", "Referee 2", "Scorer", "Timer", "24 sec operator"]

                for s_idx, s_row in skp_df_sync.iterrows():
                    m_date = str(s_row.get("Date", ""))
                    m_time = s_row.get("Time")
                    d_norm = normalize_date_str(m_date)
                    t_norm = normalize_time_str(m_time)
                    h_team = str(s_row.get("Home Team", ""))
                    a_team = str(s_row.get("Away Team", ""))
                    div_num = determine_division_for_team(h_team, tantalus_div_map)

                    for t_col in task_cols:
                        l_col = LOCK_MAP[t_col]
                        is_locked = bool(s_row.get(l_col, False))
                        assigned_player = str(s_row.get(t_col, "")).strip()

                        if not is_locked and assigned_player in valid_pool:
                            p_info = valid_pool[assigned_player]
                            if p_info["Extra points"] >= 12.0:
                                skp_df_sync.at[s_idx, t_col] = ""
                                relieved_actions.append(f"{assigned_player} ({t_col} op {d_norm})")

                                is_ref = "Referee" in t_col
                                candidates = []

                                for cand_name, cand in valid_pool.items():
                                    if cand_name == assigned_player:
                                        continue
                                    if is_ref and cand["Diploma_Rank"] == 0:
                                        continue
                                    if not is_ref and (cand["Diploma_Rank"] > 0 or cand.get("is_assistant_coach", False)):
                                        continue

                                    if is_ref and cand["Diploma"] == "BS3" and div_num != 2:
                                        continue

                                    if is_ref and "aurelie" in cand_name.lower():
                                        if not ("mse" in h_team.lower() or "mse" in a_team.lower()) or div_num != 2:
                                            continue

                                    if not is_player_eligible_for_date(cand["Season_Type"], m_date):
                                        continue

                                    if is_player_playing(cand["Team"], m_date, m_time, h_team, a_team, busy_slots):
                                        continue

                                    already_busy = False
                                    for _, check_row in skp_df_sync.iterrows():
                                        if normalize_date_str(check_row.get("Date")) == d_norm and normalize_time_str(check_row.get("Time")) == t_norm:
                                            if any(str(check_row.get(c, "")).strip() == cand_name for c in task_cols):
                                                already_busy = True
                                                break
                                    if already_busy:
                                        continue

                                    priority = 1
                                    if is_ref:
                                        dip = cand["Diploma"]
                                        if div_num in [4, 5] and dip in ["BS1", "BS2"]:
                                            priority = 10
                                        elif div_num == 3 and dip in ["BS2", "L3", "L4"]:
                                            priority = 10
                                        elif div_num <= 2 and dip in ["BS3", "L3", "L4"]:
                                            priority = 10

                                    asst_penalty = 1 if cand.get("is_assistant_coach", False) else 0
                                    candidates.append((asst_penalty, cand["Extra points"], priority, cand_name))

                                candidates.sort(key=lambda x: (x[0], x[1], -x[2]))
                                if candidates:
                                    replacement = candidates[0][3]
                                    skp_df_sync.at[s_idx, t_col] = replacement

                sheets[skp_key] = make_arrow_compatible(skp_df_sync)
                sheets = update_player_stats(sheets)

                st.session_state["changelog"].append({
                    "Actie": "Commissies Opgeslagen & Rooster Gesynchroniseerd",
                    "Details": f"Rooster direct herverdeeld. Ontlast en vervangen: {', '.join(relieved_actions) if relieved_actions else 'Geen actieve taken hoeven wijzigen'}.",
                })
                st.session_state["action_feedback"] = ("success", "Commissies zijn opgeslagen en het rooster is direct gesynchroniseerd!")
                st.rerun()

    # --- DOWNLOAD & OVERZICHT ---
    st.sidebar.divider()
    st.sidebar.header("💾 Opslaan & Downloaden")

    wb_download = openpyxl.load_workbook(io.BytesIO(st.session_state["file_bytes"]))
    for ws in wb_download.worksheets:
        if hasattr(ws, "tables"):
            ws.tables.clear()

    for sheet_name, df in sheets.items():
        if sheet_name in wb_download.sheetnames:
            ws = wb_download[sheet_name]
            clean_df = df.copy()
            for col_l in LOCK_COLS:
                if col_l in clean_df.columns:
                    clean_df = clean_df.drop(columns=[col_l])

            clean_df = clean_df.astype(object).where(pd.notna(clean_df), None)
            for r_idx, row_data in enumerate(clean_df.values, start=2):
                for c_idx, val in enumerate(row_data, start=1):
                    cell = ws.cell(row=r_idx, column=c_idx)
                    if not isinstance(cell, MergedCell):
                        cell.value = None if (pd.isna(val) or val == "") else val

    output_buffer = io.BytesIO()
    wb_download.save(output_buffer)
    output_buffer.seek(0)

    st.sidebar.download_button(
        label="📥 Download Geüpdatet Excel",
        data=output_buffer,
        file_name="SKP_Schedule_Updated.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    if st.session_state.get("indeling_gedaan", False) or st.session_state["changelog"]:
        st.divider()
        st.header("📊 Vergelijkings- en Wijzigingsweergave")

        tab_orig, tab_updated, tab_log = st.tabs(
            [
                "📁 Origineel Bestand",
                "✨ Geüpdatet Bestand",
                "📜 Logboek van Wijzigingen",
            ]
        )

        with tab_orig:
            st.info("Dit is de staat van het bestand bij aanvang.")
            orig_tab_select = st.selectbox(
                "Kies sheet (Origineel)",
                list(st.session_state["original_sheets"].keys()),
                key="orig_select",
            )
            st.dataframe(
                make_arrow_compatible(st.session_state["original_sheets"][orig_tab_select]),
                width="stretch",
            )

        with tab_updated:
            st.success("Dit is de actuele staat na de automatische indeling.")
            up_tab_select = st.selectbox(
                "Kies sheet (Geüpdatet)", list(sheets.keys()), key="up_select"
            )
            st.dataframe(make_arrow_compatible(sheets[up_tab_select]), width="stretch")

        with tab_log:
            st.warning("Chronologisch logboek van uitgevoerde acties:")
            if st.session_state["changelog"]:
                log_df = pd.DataFrame(st.session_state["changelog"])
                st.dataframe(make_arrow_compatible(log_df), width="stretch")
            else:
                st.write("Nog geen wijzigingen gelogd.")

else:
    st.warning(
        "👈 Upload aan de linkerkant je Excel-bestand om te beginnen ('SKP Schedule 26_27 (BOARD).xlsx')."
    )
